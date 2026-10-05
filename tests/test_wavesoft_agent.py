import re

import pytest

from wavesoft_agent.catalog import Catalog, Column
from wavesoft_agent.cli import main
from wavesoft_agent.connection import ConnectionSettings
from wavesoft_agent.discovery import QUERIES, discover, group_by_prefix
from wavesoft_agent.export import to_excel, to_markdown

# Réponses simulées des vues système, par nom de requête (-- q:<nom>).
ROWS = {
    "server": [{"server_name": "SRV\\WAVESOFT", "database_name": "SOCIETE1",
                "version": "Microsoft SQL Server 2019 (RTM)\nCopyright", "collation": "French_CI_AS",
                "compatibility_level": 150}],
    "objects": [
        {"object_id": 1, "schema_name": "dbo", "name": "ART_ARTICLES", "type": "U ", "created": None,
         "modified": None, "description": "Articles"},
        {"object_id": 2, "schema_name": "dbo", "name": "ART_FAMILLES", "type": "U", "created": None,
         "modified": None, "description": None},
        {"object_id": 3, "schema_name": "dbo", "name": "V_ARTICLES", "type": "V", "created": None,
         "modified": None, "description": None},
        {"object_id": 4, "schema_name": "dbo", "name": "PS_IMPORT_ARTICLE", "type": "P", "created": None,
         "modified": None, "description": None},
        {"object_id": 5, "schema_name": "dbo", "name": "TR_ART_MAJ", "type": "TR", "created": None,
         "modified": None, "description": None},
    ],
    "columns": [
        {"object_id": 1, "column_id": 1, "name": "ART_ID", "type_name": "int", "max_length": 4, "precision": 10,
         "scale": 0, "is_nullable": False, "is_identity": True, "is_computed": False, "default_definition": None,
         "computed_definition": None, "collation_name": None, "description": None},
        {"object_id": 1, "column_id": 2, "name": "ART_CODE", "type_name": "nvarchar", "max_length": 60,
         "precision": 0, "scale": 0, "is_nullable": False, "is_identity": False, "is_computed": False,
         "default_definition": None, "computed_definition": None, "collation_name": "French_CI_AS",
         "description": "Code article"},
        {"object_id": 1, "column_id": 3, "name": "ART_PRIX", "type_name": "decimal", "max_length": 9,
         "precision": 18, "scale": 4, "is_nullable": False, "is_identity": False, "is_computed": False,
         "default_definition": "((0))", "computed_definition": None, "collation_name": None, "description": None},
        {"object_id": 1, "column_id": 4, "name": "FAM_CODE", "type_name": "nvarchar", "max_length": 20,
         "precision": 0, "scale": 0, "is_nullable": True, "is_identity": False, "is_computed": False,
         "default_definition": None, "computed_definition": None, "collation_name": None, "description": None},
        {"object_id": 2, "column_id": 1, "name": "FAM_CODE", "type_name": "nvarchar", "max_length": 20,
         "precision": 0, "scale": 0, "is_nullable": False, "is_identity": False, "is_computed": False,
         "default_definition": None, "computed_definition": None, "collation_name": None, "description": None},
    ],
    "indexes": [
        {"object_id": 1, "name": "PK_ART", "is_primary_key": True, "is_unique": True, "type_desc": "CLUSTERED",
         "column_name": "ART_ID"},
        {"object_id": 1, "name": "UQ_ART_CODE", "is_primary_key": False, "is_unique": True,
         "type_desc": "NONCLUSTERED", "column_name": "ART_CODE"},
    ],
    "foreign_keys": [
        {"object_id": 1, "name": "FK_ART_FAM", "column_name": "FAM_CODE", "ref_schema": "dbo",
         "ref_table": "ART_FAMILLES", "ref_column": "FAM_CODE", "on_delete": "NO_ACTION",
         "on_update": "NO_ACTION", "is_disabled": False},
    ],
    "checks": [{"object_id": 1, "name": "CK_PRIX", "definition": "([ART_PRIX]>=(0))"}],
    "parameters": [
        {"object_id": 4, "parameter_id": 1, "name": "@CODE", "type_name": "nvarchar", "max_length": 60,
         "precision": 0, "scale": 0, "is_output": False, "has_default_value": False},
        {"object_id": 4, "parameter_id": 2, "name": "@ID", "type_name": "int", "max_length": 4,
         "precision": 10, "scale": 0, "is_output": True, "has_default_value": False},
    ],
    "definitions": [{"object_id": 4, "definition": "CREATE PROCEDURE PS_IMPORT_ARTICLE ... INSERT ART_ARTICLES"}],
    "triggers": [
        {"object_id": 5, "parent": "dbo.ART_ARTICLES", "is_disabled": False, "event": "INSERT"},
        {"object_id": 5, "parent": "dbo.ART_ARTICLES", "is_disabled": False, "event": "UPDATE"},
    ],
    "row_counts": [{"object_id": 1, "row_count": 1234}],
    "synonyms": [],
    "dependencies": [
        {"object_id": 3, "referenced": "dbo.ART_ARTICLES"},
        {"object_id": 4, "referenced": "dbo.ART_ARTICLES"},
    ],
}


class FakeCursor:
    def __init__(self, rows, fail):
        self._rows, self._fail, self._result = rows, fail, []
        self.description = None

    def execute(self, sql):
        name = re.match(r"-- q:(\w+)", sql).group(1)
        if name in self._fail:
            raise PermissionError("VIEW DEFINITION refusé")
        rows = self._rows.get(name, [])
        keys = list(rows[0]) if rows else ["x"]
        self.description = [(k,) for k in keys]
        self._result = [tuple(r[k] for k in keys) for r in rows]

    def fetchall(self):
        return self._result

    def close(self):
        pass


class FakeConn:
    def __init__(self, rows=ROWS, fail=()):
        self.rows, self.fail = rows, set(fail)

    def cursor(self):
        return FakeCursor(self.rows, self.fail)


@pytest.fixture
def catalog():
    return discover(FakeConn())


def test_every_query_is_tagged():
    for name, sql in QUERIES.items():
        assert sql.startswith(f"-- q:{name}\n")
        assert not re.search(r"\b(INSERT|UPDATE|DELETE|MERGE|EXEC|DROP|ALTER|CREATE)\b", sql.split("\n", 1)[1])


def test_discover_builds_full_catalog(catalog):
    assert catalog.database == "SOCIETE1"
    assert catalog.server_version == "Microsoft SQL Server 2019 (RTM)"
    assert catalog.counts() == {"procedure": 1, "table": 2, "trigger": 1, "view": 1}

    art = catalog.get("art_articles")
    assert art.row_count == 1234
    assert art.primary_key.columns == ["ART_ID"] and art.primary_key.clustered
    assert [k.kind for k in art.keys] == ["primary_key", "unique"]
    assert art.foreign_keys[0].ref_table == "ART_FAMILLES"
    assert art.checks == {"CK_PRIX": "([ART_PRIX]>=(0))"}
    assert art.column("art_code").sql_type == "nvarchar(30)"
    assert art.column("ART_PRIX").sql_type == "decimal(18,4)"
    # identité, défaut et nullable ne sont pas à fournir dans un fichier d'intégration
    assert [c.name for c in art.required_columns()] == ["ART_CODE"]

    proc = catalog.get("dbo.PS_IMPORT_ARTICLE")
    assert [(p.name, p.output) for p in proc.parameters] == [("@CODE", False), ("@ID", True)]
    assert "INSERT ART_ARTICLES" in proc.definition

    trig = catalog.get("TR_ART_MAJ")
    assert trig.parent == "dbo.ART_ARTICLES" and trig.trigger_events == ["INSERT", "UPDATE"]

    assert {o.name for o in catalog.referenced_by("ART_ARTICLES")} == {"V_ARTICLES", "PS_IMPORT_ARTICLE"}


def test_missing_view_definition_right_is_a_warning_not_a_failure():
    warnings = []
    cat = discover(FakeConn(fail={"definitions"}), warnings)
    assert cat.get("PS_IMPORT_ARTICLE").definition is None
    assert any("VIEW DEFINITION" in w for w in warnings)


def test_mandatory_query_failure_raises():
    with pytest.raises(PermissionError):
        discover(FakeConn(fail={"objects"}))


def test_json_round_trip(catalog, tmp_path):
    path = catalog.save(tmp_path / "catalog.json")
    again = Catalog.load(path)
    assert again.to_dict() == catalog.to_dict()
    assert isinstance(again.get("ART_ARTICLES").columns[0], Column)


def test_exports(catalog, tmp_path):
    md = to_markdown(catalog)
    assert "### dbo.ART_ARTICLES" in md and "| 2 | ART_CODE | nvarchar(30) |" in md
    from openpyxl import load_workbook

    wb = load_workbook(to_excel(catalog, tmp_path / "c.xlsx"))
    assert wb.sheetnames == ["Objets", "Colonnes", "Paramètres"]
    assert wb["Colonnes"].max_row == 1 + 5


def test_group_by_prefix(catalog):
    assert group_by_prefix(catalog)["ART"] == ["dbo.ART_ARTICLES", "dbo.ART_FAMILLES"]


def test_cli_search_and_describe(catalog, tmp_path, capsys):
    path = str(catalog.save(tmp_path / "catalog.json"))
    assert main(["search", path, "^FAM_"]) == 0
    out = capsys.readouterr().out
    assert "dbo.ART_ARTICLES.FAM_CODE" in out and "dbo.ART_FAMILLES.FAM_CODE" in out
    assert main(["describe", path, "ART_ARTICLES"]) == 0
    assert "Utilisé par" in capsys.readouterr().out
    assert main(["search", path, "INTROUVABLE"]) == 1


def test_connection_string(tmp_path, monkeypatch):
    for k in ("WAVESOFT_SERVER", "WAVESOFT_DATABASE", "WAVESOFT_USER", "WAVESOFT_PASSWORD"):
        monkeypatch.delenv(k, raising=False)
    env = tmp_path / ".env"
    env.write_text("WAVESOFT_SERVER=SRV\\WS\nWAVESOFT_DATABASE=SOC1\nWAVESOFT_USER=agent\nWAVESOFT_PASSWORD=a;b}c\n")
    cs = ConnectionSettings.from_env(env).connection_string()
    assert "SERVER=SRV\\WS;" in cs and "ApplicationIntent=ReadOnly" in cs and "PWD={a;b}}c};" in cs
    trusted = ConnectionSettings.from_env(None, server="S", database="D").connection_string()
    assert "Trusted_Connection=yes" in trusted
    with pytest.raises(ValueError):
        ConnectionSettings.from_env(None)
