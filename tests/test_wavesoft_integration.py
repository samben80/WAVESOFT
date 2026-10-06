"""Tests de l'intégration directe, sur un catalogue et une base simulés."""

import re
from datetime import date
from decimal import Decimal

import pytest

from wavesoft_agent.catalog import Catalog, Column, DbObject
from wavesoft_agent.integration import (Config, convertir, ecrire_modele, entetes, integrer, preparer,
                                        verifier_config)


def col(name, t="varchar", n=None, nullable=True, **kw):
    return Column(name, 0, t, max_length=n, nullable=nullable, **kw)


def table(name, *cols):
    return DbObject("dbo", name, "U", "table", columns=list(cols))


def catalogue():
    return Catalog("srv", "DOSSIER", "2026-10-06", objects=[
        table("FAMILLES", col("FAMID", "int", nullable=False), col("FAMCODE", n=10, nullable=False),
              col("FAMINTITULE", n=50)),
        table("FAMILLES_P", col("FAMID", "int", nullable=False)),
        table("ARTICLES", col("ARTID", "int", nullable=False), col("ARTCODE", n=25, nullable=False),
              col("ARTDESIGNATION", n=60, nullable=False), col("FAMID", "int", nullable=False),
              col("ARTPRIXVENTE", "MNTCPT", base_type="decimal", precision=18, scale=4),
              col("ARTISACTIF", "BOOL", base_type="char", n=1)),
        table("ARTICLES_P", col("ARTID", "int", nullable=False)),
        table("TIERS", col("TIRID", "int", nullable=False), col("TIRCODE", n=17, nullable=False),
              col("TIRTYPE", "char", n=1, nullable=False), col("TIRSOCIETE", n=60)),
        table("PIECEVENTES", col("PCVID", "int", nullable=False), col("PCVNUMEXT", n=20),
              col("PINCODE", n=10, nullable=False), col("TIRID", "int", nullable=False),
              col("PCVDATEEFFET", "datetime", nullable=False)),
        table("PIECEVENTELIGNES", col("PLVID", "int", nullable=False, identity=True),
              col("PCVID", "int", nullable=False), col("ARTID", "int", nullable=False),
              col("PLVQTE", "decimal", precision=18, scale=6, nullable=False)),
        table("PIECENATURES", col("PINCODE", n=10), col("PININTITULE", n=50)),
    ])


class FakeDb:
    """Base en mémoire qui comprend les requêtes émises par integration.Base."""

    def __init__(self, existing=None):
        self.rows = {k: list(v) for k, v in (existing or {}).items()}
        self.ids = {}
        self.committed = self.rolled_back = False
        self.last = None

    def cursor(self):
        return self

    def execute(self, sql, *params):
        if "ws_sp_GetIdTable" in sql:
            self.ids[params[0]] = self.ids.get(params[0], 100) + 1
            self.last = [(self.ids[params[0]],)]
        elif sql.startswith("INSERT"):
            t = re.search(r"dbo\.\[(\w+)\]", sql).group(1)
            cols = re.findall(r"\[(\w+)\]", sql.split("VALUES")[0])[1:]
            row = dict(zip(cols, params))
            if t == "PIECEVENTELIGNES":
                row["PLVID"] = len(self.rows.get(t, [])) + 1
            self.rows.setdefault(t, []).append(row)
        elif "SCOPE_IDENTITY" in sql:
            self.last = [(len(self.rows["PIECEVENTELIGNES"]),)]
        elif "ORDER BY 1" in sql:
            self.last = [(r["PINCODE"], r.get("PININTITULE")) for r in self.rows.get("PIECENATURES", [])]
        elif sql.startswith("SELECT ["):
            m = re.match(r"SELECT \[(\w+)\] FROM dbo\.\[(\w+)\] WHERE (.*)", sql)
            id_col, t, where = m.groups()
            keys = re.findall(r"\[(\w+)\] = \?", where)
            self.last = [(r[id_col],) for r in self.rows.get(t, [])
                         if all(str(r.get(k)) == str(v) for k, v in zip(keys, params))][:1]
        else:
            raise AssertionError(sql)

    def fetchall(self):
        return self.last

    def fetchone(self):
        return self.last[0] if self.last else None

    def commit(self):
        self.committed = True

    def rollback(self):
        self.rolled_back = True


@pytest.fixture
def config():
    return Config.load()


def remplir(path, feuilles):
    from openpyxl import load_workbook

    wb = load_workbook(path)
    for nom, lignes in feuilles.items():
        ws = wb[nom]
        hdr = [c.value for c in ws[1]]
        for i, valeurs in enumerate(lignes, start=3):
            for k, v in valeurs.items():
                ws.cell(row=i, column=hdr.index(k) + 1, value=v)
    wb.save(path)
    return path


@pytest.fixture
def classeur(tmp_path, config):
    p = ecrire_modele(config, catalogue(), tmp_path / "modele.xlsx")
    return remplir(p, {
        "Familles": [{"FAMCODE": "VELO", "FAMINTITULE": "Vélos"}],
        "Articles": [
            {"ARTCODE": "A1", "ARTDESIGNATION": "Vélo", "FAMILLE": "VELO", "ARTPRIXVENTE": 199.9, "ARTISACTIF": "oui"},
            {"ARTCODE": 2.0, "ARTDESIGNATION": "Casque"},  # famille vide -> DEFAULT ; code saisi en nombre
        ],
        "Clients": [{"TIRCODE": "C1", "TIRSOCIETE": "SA Raymond"}],
        "Ventes": [{"PIECE": "P1", "CLIENT": "C1", "PCVDATEEFFET": date(2026, 1, 5), "PCVNUMEXT": "H001"}],
        "Ventes lignes": [{"PIECE": "P1", "ARTICLE": "A1", "PLVQTE": 2}, {"PIECE": "P1", "ARTICLE": "2", "PLVQTE": 1}],
    })


def test_config_checked_against_catalog(config):
    issues = [str(i) for i in verifier_config(config, catalogue())]
    assert any("table PRODUITS absente" in i for i in issues)
    assert any("table PIECEACHATS absente" in i for i in issues)
    assert not any("FAMILLES" in i or "ARTICLES" in i for i in issues)


def test_template_uses_real_columns(config, tmp_path):
    cat = catalogue()
    art = entetes(config.etape("articles"), cat.get("ARTICLES"))
    assert art[:3] == ["FAMILLE", "ARTCODE", "ARTDESIGNATION"]  # référence, puis obligatoires
    assert "ARTID" not in art and "FAMID" not in art
    lignes = entetes(config.etape("ventes_lignes"), cat.get("PIECEVENTELIGNES"))
    assert lignes == ["PIECE", "ARTICLE", "PLVQTE"]  # ni PLVID (identité) ni PCVID (parent)
    from openpyxl import load_workbook

    wb = load_workbook(ecrire_modele(config, cat, tmp_path / "m.xlsx"))
    assert wb.sheetnames[:3] == ["Lisez-moi", "Familles", "Articles"] and "Produits" not in wb.sheetnames
    assert "vide = DEFAULT" in wb["Articles"]["A2"].value


def test_value_conversion():
    assert convertir(col("C", n=3), 12.0) == "12"
    with pytest.raises(ValueError, match="maximum 3"):
        convertir(col("C", n=3), "abcd")
    assert convertir(col("N", "nvarchar", n=6), "abc") == "abc"  # nvarchar : longueur en octets / 2
    assert convertir(col("B", "BOOL", base_type="char", n=1), True) == "O"
    assert convertir(col("M", "decimal", precision=6, scale=2), "1 234,5") == Decimal("1234.50")
    with pytest.raises(ValueError, match="trop grand"):
        convertir(col("M", "decimal", precision=4, scale=2), 1000)
    assert convertir(col("D", "datetime"), "05/01/2026").day == 5
    with pytest.raises(ValueError, match="entier"):
        convertir(col("I", "int"), 1.5)


def test_simulation_writes_in_order_then_rolls_back(config, classeur):
    cat = catalogue()
    plan = preparer(config, cat, classeur, {"vente": "FACHISTO"})
    assert plan.erreurs == []
    db = FakeDb({"FAMILLES": [{"FAMID": 1, "FAMCODE": "DEFAULT"}]})
    bilan = integrer(plan, config, cat, db)
    assert db.rolled_back and not db.committed and not bilan.execute
    assert bilan.crees == {"familles": 1, "articles": 2, "clients": 1, "ventes": 1, "ventes_lignes": 2}
    velo = db.rows["FAMILLES"][1]
    assert velo["FAMID"] == 101 and db.rows["FAMILLES_P"] == [{"FAMID": 101}]
    a1, a2 = db.rows["ARTICLES"]
    assert a1["FAMID"] == 101 and a2["FAMID"] == 1 and a2["ARTCODE"] == "2"
    assert a1["ARTPRIXVENTE"] == Decimal("199.9000") and a1["ARTISACTIF"] == "O"
    assert db.rows["TIERS"][0]["TIRTYPE"] == "C"
    piece = db.rows["PIECEVENTES"][0]
    assert piece["PINCODE"] == "FACHISTO" and piece["TIRID"] == db.rows["TIERS"][0]["TIRID"]
    assert [l["PCVID"] for l in db.rows["PIECEVENTELIGNES"]] == [piece["PCVID"]] * 2
    assert [l["ARTID"] for l in db.rows["PIECEVENTELIGNES"]] == [a1["ARTID"], a2["ARTID"]]
    assert "PLVID" not in db.rows["PIECEVENTELIGNES"][0] or db.rows["PIECEVENTELIGNES"][0]["PLVID"] == 1
    assert any("ne sont pas recalculés" in str(a) for a in bilan.anomalies)


def test_execute_commits_and_skips_existing(config, classeur):
    cat = catalogue()
    plan = preparer(config, cat, classeur, {"vente": "FACHISTO"})
    db = FakeDb({"FAMILLES": [{"FAMID": 1, "FAMCODE": "DEFAULT"}],
                 "TIERS": [{"TIRID": 7, "TIRCODE": "C1", "TIRTYPE": "C"}],
                 "PIECEVENTES": [{"PCVID": 9, "PCVNUMEXT": "H001"}]})
    bilan = integrer(plan, config, cat, db, executer=True)
    assert db.committed and bilan.execute
    assert bilan.existants == {"clients": 1, "ventes": 1}
    assert "PIECEVENTELIGNES" not in db.rows  # jamais de lignes ajoutées à une pièce existante


def test_unknown_reference_rolls_everything_back(config, tmp_path):
    cat = catalogue()
    p = remplir(ecrire_modele(config, cat, tmp_path / "m.xlsx"),
                {"Articles": [{"ARTCODE": "A1", "ARTDESIGNATION": "x", "FAMILLE": "INCONNUE"}]})
    db = FakeDb({"FAMILLES": [{"FAMID": 1, "FAMCODE": "DEFAULT"}]})
    bilan = integrer(preparer(config, cat, p), config, cat, db, executer=True)
    assert any("FAMILLE INCONNUE introuvable" in str(a) for a in bilan.anomalies)
    assert db.rolled_back and not db.committed


def test_workbook_errors_block_everything(config, tmp_path):
    cat = catalogue()
    p = remplir(ecrire_modele(config, cat, tmp_path / "m.xlsx"), {
        "Articles": [{"ARTCODE": "A1"}, {"ARTCODE": "A1", "ARTDESIGNATION": "x" * 61}],
        "Ventes": [{"PIECE": "P1", "CLIENT": "C1", "PCVDATEEFFET": "hier"}],
        "Ventes lignes": [{"PIECE": "P2", "ARTICLE": "A1", "PLVQTE": 1}],
    })
    plan = preparer(config, cat, p)
    msgs = [str(a) for a in plan.erreurs]
    assert any("Articles ligne 3 : ARTDESIGNATION : obligatoire" in m for m in msgs)
    assert any("61 caractères, maximum 60" in m for m in msgs)
    assert any("ARTCODE A1 en double" in m for m in msgs)
    assert any("nature de pièce de vente non choisie" in m for m in msgs)
    assert any("PCVDATEEFFET : date attendue" in m for m in msgs)
    assert any("PIECE P2 sans entête" in m for m in msgs)
    with pytest.raises(ValueError, match="rien n'est écrit"):
        integrer(plan, config, cat, FakeDb())


def test_cli_integrer(config, classeur, tmp_path, monkeypatch, capsys):
    import contextlib

    from wavesoft_agent import cli

    cat_path = catalogue().save(tmp_path / "catalog.json")
    db = FakeDb({"FAMILLES": [{"FAMID": 1, "FAMCODE": "DEFAULT"}], "PIECENATURES": [{"PINCODE": "FACHISTO"}]})
    monkeypatch.setattr(cli, "_connect", lambda args, write=False: contextlib.nullcontext(db))
    assert cli.main(["integrer", str(cat_path), str(classeur), "--nature-vente", "fachisto"]) == 0
    out = capsys.readouterr().out
    assert "Simulation" in out and db.rolled_back and not db.committed
    assert db.rows["PIECEVENTES"][0]["PINCODE"] == "FACHISTO"
    assert cli.main(["integrer", str(cat_path), str(classeur), "--nature-vente", "X"]) == 2  # nature inconnue
    assert "nature X inconnue" in capsys.readouterr().err
    assert cli.main(["integrer", str(cat_path), str(classeur), "--nature-vente", "FACHISTO", "--executer"]) == 2
    assert not db.committed  # --executer sans terminal ni --oui : refusé
    assert cli.main(["integration-verifier", str(cat_path)]) == 1
    assert "table PRODUITS absente" in capsys.readouterr().out
