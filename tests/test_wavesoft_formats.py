"""Tests du moteur de formats, sur une définition fictive (les vraies sont privées)."""

import json
from datetime import date

import pytest

from wavesoft_agent.cli import main
from wavesoft_agent.excel_io import build_file, write_template
from wavesoft_agent.formats import FormatSpec, check_value, format_value, load_specs, FieldSpec
from wavesoft_agent.importfile import Document, check_file, split_records, write_file

SPEC = {
    "code": "TEST",
    "titre": "Pièces de test",
    "enregistrements": {
        "E": {"niveau": "entete", "unique": True, "champs": [
            {"n": 1, "nom": "type_ligne", "libelle": "Type"},
            {"n": 2, "nom": "nature", "libelle": "Nature", "obligatoire": True},
            {"n": 3, "nom": "date_effet", "libelle": "Date d'effet", "type": "date", "obligatoire": True},
            {"n": 4, "nom": "client", "libelle": "Code client", "max": 17, "obligatoire": True},
            {"n": 5, "nom": "vide_5", "type": "vide"},
            {"n": 6, "nom": "ht", "libelle": "H.T.", "type": "on"},
        ]},
        "ED": {"niveau": "entete", "repetition_depuis": 2, "champs": [
            {"n": 1, "nom": "type_ligne"},
            {"n": 2, "nom": "code_champ", "max": 25},
            {"n": 3, "nom": "valeur", "type": "libre"},
        ]},
        "NO": {"niveau": "entete", "champs": [
            {"n": 1, "nom": "type_ligne"}, {"n": 2, "nom": "note", "type": "texte_long"},
        ]},
        "LA": {"niveau": "ligne", "champs": [
            {"n": 1, "nom": "type_ligne"},
            {"n": 2, "nom": "article", "libelle": "Code article", "max": 25, "obligatoire": True},
            {"n": 3, "nom": "quantite", "libelle": "Quantité", "type": "decimal", "decimales": 6, "obligatoire": True},
            {"n": 4, "nom": "remise_type", "type": "liste", "valeurs": ["P", "M"]},
        ]},
        "LD": {"niveau": "ligne", "suit": ["LA"], "repetition_depuis": 2, "champs": [
            {"n": 1, "nom": "type_ligne"}, {"n": 2, "nom": "code_champ"}, {"n": 3, "nom": "valeur", "type": "libre"},
        ]},
        "LT": {"niveau": "ligne", "champs": [{"n": 1, "nom": "type_ligne"}, {"n": 2, "nom": "designation"}]},
        "LC": {"comme": "LT", "libelle": "Commentaire"},
        "CLO": {"niveau": "commande", "champs": [
            {"n": 1, "nom": "type_ligne"}, {"n": 2, "nom": "ref_piece", "obligatoire": True},
        ]},
    },
}


@pytest.fixture
def spec():
    return FormatSpec.from_dict(SPEC)


@pytest.fixture
def spec_dir(tmp_path):
    d = tmp_path / "formats"
    d.mkdir()
    (d / "TEST.json").write_text(json.dumps(SPEC), encoding="utf-8")
    return d


def write(tmp_path, text, name="f.txt"):
    p = tmp_path / name
    p.write_bytes(text.encode("cp1252"))
    return p


def test_alias_record_takes_its_own_code(spec):
    assert spec.enregistrements["LC"].champs[0].valeurs == ["LC"]
    assert spec.enregistrements["LA"].champs[0].valeurs == ["LA"]


def test_value_formatting_and_checks():
    dec = FieldSpec(1, "q", type="decimal", decimales=6)
    assert format_value(dec, 2) == "2.000000"
    assert format_value(FieldSpec(1, "d", type="date"), date(2026, 1, 5)) == "05/01/2026"
    assert format_value(FieldSpec(1, "c"), 2.0) == "2"  # code saisi comme nombre dans Excel
    assert format_value(FieldSpec(1, "b", type="on"), True) == "O"
    assert check_value(dec, "1,5") and check_value(FieldSpec(1, "d", type="date"), "2026-01-05")
    assert check_value(FieldSpec(1, "t", max=3), "abcd") == "4 caractères, maximum 3"
    assert check_value(FieldSpec(1, "t", obligatoire=True), "") == "obligatoire"


def test_split_records_quoted_multiline():
    recs = split_records('E;X;\r\nNO;"ligne 1\r\nligne 2; suite";\r\nLA;A;1\r\n', ";")
    assert [r.code for r in recs] == ["E", "NO", "LA"]
    assert recs[1].values == ["NO", "ligne 1\r\nligne 2; suite"]
    assert recs[2].line == 4


def test_check_valid_file(spec, tmp_path):
    f = write(tmp_path, "E;CDECLI;05/10/2026;0002;;O;\r\nED;CP1;a;CP2;b;\r\nLA;ART1;1.000000;M;\r\nLD;CP3;x;\r\n"
                        "LC;un commentaire;\r\nCLO;CC1;\r\n")
    assert check_file(f, spec) == []


def test_check_reports_each_problem(spec, tmp_path):
    f = write(tmp_path, "LA;ART;1\r\nE;CDECLI;2026-10-05;;X;Z\r\nE;CDECLI;05/10/2026;1\r\nLD;CP;x\r\n"
                        "LA;ART;1;Q;de trop\r\nZZ;?\r\n")
    msgs = [str(i) for i in check_file(f, spec)]
    assert any("ligne 1 (LA) : doit suivre une ligne E" in m for m in msgs)
    assert any("Date d'effet : date au format dd/mm/yyyy" in m for m in msgs)
    assert any("Code client : obligatoire" in m for m in msgs)
    assert any("ATTENTION ligne 2 (E, colonne 5)" in m for m in msgs)  # colonne VIDE renseignée
    assert any("H.T. : O ou N attendu" in m for m in msgs)
    assert any("ligne 4 (LD) : doit suivre une ligne LA" in m for m in msgs)
    assert any("valeur parmi P, M" in m for m in msgs)
    assert any("colonne en trop" in m for m in msgs)
    assert any("type d'enregistrement inconnu" in m for m in msgs)


def test_write_file_round_trip(spec, tmp_path):
    doc = (Document().add("E", nature="CDECLI", date_effet=date(2026, 10, 5), client="0002", ht=False)
           .add("ED", paires=[("CP1", "a"), ("CP2", 3)])
           .add("NO", note="Livrer\nle matin")
           .add("LA", article="ART1", quantite=2).add("LD", code_champ="CP3", valeur="é"))
    out = tmp_path / "out.txt"
    assert write_file(out, spec, [doc, Document().add("CLO", ref_piece="CC1")]) == []
    text = out.read_bytes().decode("cp1252")
    assert text.splitlines()[0] == "E;CDECLI;05/10/2026;0002;;N;"
    assert "ED;CP1;a;CP2;3;" in text and 'NO;"Livrer\nle matin";' in text and "LA;ART1;2.000000;" in text
    assert check_file(out, spec) == []


def test_write_file_refuses_bad_data(spec, tmp_path):
    out = tmp_path / "out.txt"
    bad = Document().add("E", nature="CDECLI", date_effet="05/10/2026", client="X;Y").add("LA", article="A", quantite=1)
    issues = write_file(out, spec, [bad])
    assert any("séparateur" in i.message for i in issues) and not out.exists()
    issues = write_file(out, spec, [Document().add("E", nature="N", date_effet="05/10/2026", client="€uro").add(
        "NO", note='un "guillemet"')])
    assert any("guillemet interdit" in i.message for i in issues) and not out.exists()
    issues = write_file(out, spec, [Document().add("E", nature="N", date_effet="05/10/2026", client="漢")])
    assert any("impossible à écrire en cp1252" in i.message for i in issues) and not out.exists()


def test_excel_template_and_build(spec, tmp_path):
    from openpyxl import load_workbook

    tpl = write_template(spec, tmp_path / "modele.xlsx")
    wb = load_workbook(tpl)
    assert wb.sheetnames[0] == "Lisez-moi" and "LA" in wb.sheetnames
    assert [c.value for c in wb["LA"][1]] == ["PIECE", "ORDRE", "article", "quantite", "remise_type"]
    assert [c.value for c in wb["E"][1]] == ["PIECE", "nature", "date_effet", "client", "ht"]  # VIDE masquée

    def put(sheet, **kv):
        ws = wb[sheet]
        hdr = [c.value for c in ws[1]]
        row = 3
        while ws.cell(row=row, column=1).value is not None:
            row += 1
        for k, v in kv.items():
            ws.cell(row=row, column=hdr.index(k) + 1, value=v)

    put("E", PIECE="P1", nature="CDECLI", date_effet=date(2026, 10, 5), client="0002")
    put("ED", PIECE="P1", code_champ="CP1", valeur="a")
    put("ED", PIECE="P1", code_champ="CP2", valeur="b")
    put("LA", PIECE="P1", ORDRE=2, article="B", quantite=1)
    put("LD", PIECE="P1", ORDRE=1, code_champ="CPL", valeur="x")
    put("LA", PIECE="P1", ORDRE=1, article="A", quantite=2.5)
    put("LC", PIECE="P1", ORDRE=0, designation="Titre")
    put("CLO", ORDRE=1, ref_piece="CC9")
    filled = tmp_path / "rempli.xlsx"
    wb.save(filled)

    out = tmp_path / "import.txt"
    assert build_file(spec, filled, out) == []
    assert out.read_text(encoding="cp1252").splitlines() == [
        "E;CDECLI;05/10/2026;0002;", "ED;CP1;a;CP2;b;", "LC;Titre;", "LA;A;2.500000;", "LD;CPL;x;",
        "LA;B;1.000000;", "CLO;CC9;",
    ]

    put("LA", PIECE="P2", ORDRE=1, article="C", quantite=1)
    put("LA", PIECE="P1", ORDRE=3, article="D", quantite="beaucoup")
    wb.save(filled)
    out.unlink()
    msgs = [i.message for i in build_file(spec, filled, out)]
    assert any("PIECE 'P2' sans entête E" in m for m in msgs)
    assert any(m.startswith("feuille LA ligne 6 : Quantité") for m in msgs)
    assert not out.exists()


def test_cli_formats(spec_dir, tmp_path, capsys):
    assert main(["formats", "--specs", str(spec_dir)]) == 0
    assert "TEST" in capsys.readouterr().out
    f = write(tmp_path, "E;N;05/10/2026;C1\r\n")
    assert main(["check-file", "test", str(f), "--specs", str(spec_dir)]) == 0
    assert main(["template", "TEST", "--out", str(tmp_path / "m.xlsx"), "--records", "E,LA",
                 "--specs", str(spec_dir)]) == 0
    assert main(["check-file", "AUTRE", str(f), "--specs", str(spec_dir)]) == 2
    assert main(["formats", "--specs", str(tmp_path / "absent")]) == 2


def test_code_less_format(tmp_path):
    spec = FormatSpec.from_dict({"code": "L", "titre": "Lignes", "documents": {"debut": None}, "enregistrements": {
        "*": {"niveau": "ligne", "champs": [
            {"n": 1, "nom": "article", "max": 25, "obligatoire": True},
            {"n": 2, "nom": "quantite", "type": "decimal", "decimales": 6, "obligatoire": True}]}}})
    assert spec.sans_code
    f = write(tmp_path, "ART1;2\r\nART2;x\r\n")
    issues = check_file(f, spec)
    assert len(issues) == 1 and issues[0].line == 2
    out = tmp_path / "o.txt"
    assert write_file(out, spec, [Document().add("*", article="A", quantite=1)]) == []
    assert out.read_bytes() == b"A;1.000000;\r\n"


def test_load_specs_sorted(spec_dir):
    assert list(load_specs(spec_dir)) == ["TEST"]
