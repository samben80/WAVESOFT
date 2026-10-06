"""Modèle Excel à faire remplir et génération du fichier d'import à partir de ce modèle.

Le classeur contient une feuille par type d'enregistrement du format (``E``,
``AF``, ``LA``…). Ligne 1 : nom technique des colonnes (ne pas modifier) ;
ligne 2 : description ; données à partir de la ligne 3.

- ``PIECE`` relie les enregistrements d'une même pièce (n'importe quel code
  choisi par l'utilisateur, il n'est pas écrit dans le fichier).
- ``ORDRE`` donne l'ordre des lignes dans la pièce. Un ``LD``, ``AN`` ou ``LP``
  prend l'``ORDRE`` de la ligne article qu'il complète.
"""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

from .formats import FieldSpec, FormatSpec, RecordSpec
from .importfile import Document, Issue, write_file

KEY, ORDER = "PIECE", "ORDRE"
DATA_ROW = 3


def _data_fields(rec: RecordSpec, code_less: bool) -> list[FieldSpec]:
    return [f for f in rec.champs if f.type != "vide" and (code_less or f.n != 1)]


def _describe(f: FieldSpec) -> str:
    parts = [f.libelle or f.nom]
    if f.type in ("texte", "texte_long") and f.max:
        parts.append(f"texte {f.max} car.")
    elif f.type == "decimal":
        parts.append(f"nombre ({f.decimales} déc.)" if f.decimales is not None else "nombre")
    elif f.type == "on":
        parts.append("O/N")
    elif f.type == "liste" and f.valeurs:
        parts.append("/".join(f.valeurs))
    elif f.type in ("date", "heure", "entier"):
        parts.append(f.type)
    if f.obligatoire:
        parts.append("OBLIGATOIRE")
    if f.export_seulement:
        parts.append("export seulement")
    if f.note:
        parts.append(f.note)
    return " · ".join(parts)


def _columns(spec: FormatSpec, rec: RecordSpec) -> list[str]:
    if spec.sans_code:
        return []
    cols = [KEY] if rec.niveau != "commande" or spec.debut_document else []
    if rec.niveau in ("ligne", "commande"):
        cols.append(ORDER)
    return cols


def write_template(spec: FormatSpec, path: str | Path, records: list[str] | None = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    info = wb.active
    info.title = "Lisez-moi"
    info.append([f"{spec.code} — {spec.titre}" + (f" (fiche v{spec.version})" if spec.version else "")])
    info["A1"].font = Font(bold=True, size=13)
    for line in [
        "Une feuille par type d'enregistrement. Ne modifiez pas la ligne 1 (noms techniques).",
        "La ligne 2 décrit chaque colonne. Saisissez les données à partir de la ligne 3.",
        "PIECE : code libre qui regroupe les enregistrements d'une même pièce.",
        "ORDRE : ordre des lignes dans la pièce ; LD/AN/LP reprennent l'ORDRE de la ligne qu'ils complètent.",
        "Les feuilles laissées vides sont ignorées.",
        spec.note or "",
    ]:
        info.append([line])
    info.column_dimensions["A"].width = 110

    wanted = records or list(spec.enregistrements)
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    req_font = Font(bold=True, color="C00000")
    for code in wanted:
        rec = spec.enregistrements[code]
        ws = wb.create_sheet(("Lignes" if code == "*" else code)[:31])
        fields = _data_fields(rec, spec.sans_code)
        cols = _columns(spec, rec)
        ws.append(cols + [f.nom for f in fields])
        ws.append([("Code pièce (libre)" if c == KEY else "Ordre (entier)") for c in cols] + [_describe(f) for f in fields])
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = head_fill
        for cell, f in zip(ws[2][len(cols):], fields):
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if f.obligatoire:
                cell.font = req_font
        ws.row_dimensions[2].height = 60
        ws.freeze_panes = ws.cell(row=DATA_ROW, column=len(cols) + 1)
        for i, f in enumerate(fields, start=len(cols) + 1):
            width = min(max((f.max or 12) + 2, 12), 40)
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
            if f.type in ("texte", "texte_long", "liste", "on"):
                for r in range(DATA_ROW, DATA_ROW + 500):  # codes "0002" restent du texte
                    ws.cell(row=r, column=i).number_format = "@"
        ws.sheet_properties.tabColor = {"entete": "4472C4", "ligne": "70AD47", "commande": "ED7D31"}[rec.niveau]
        if rec.note:
            ws.cell(row=2, column=len(cols) + len(fields) + 2, value=rec.note).font = Font(italic=True)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


@dataclass
class _Row:
    code: str
    sheet: str
    row: int
    key: str | None
    order: float | None
    values: dict


def _read_rows(spec: FormatSpec, path: str | Path) -> list[_Row]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    rows: list[_Row] = []
    for ws in wb.worksheets:
        code = "*" if spec.sans_code and ws.title == "Lignes" else ws.title
        rec = spec.enregistrements.get(code)
        if rec is None:
            continue
        header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
        valid = {f.nom for f in rec.champs} | {KEY, ORDER}
        unknown = [h for h in header if h and h not in valid]
        if unknown:
            raise ValueError(f"Feuille {ws.title} : colonnes inconnues {unknown} (ligne 1 modifiée ?)")
        for r_idx, row in enumerate(ws.iter_rows(min_row=DATA_ROW, values_only=True), start=DATA_ROW):
            values = {h: v for h, v in zip(header, row) if h and v not in (None, "")}
            if not values:
                continue
            key = values.pop(KEY, None)
            order = values.pop(ORDER, None)
            rows.append(_Row(code, ws.title, r_idx, None if key is None else str(key).strip(),
                             None if order is None else float(order), values))
    return rows


def build_documents(spec: FormatSpec, rows: list[_Row]) -> tuple[list[Document], list[tuple[str, int]], list[Issue]]:
    """Assemble les pièces. Renvoie aussi l'origine (feuille, ligne) de chaque enregistrement."""
    issues: list[Issue] = []
    if spec.sans_code:
        docs = [Document([("*", r.values)]) for r in rows]
        return docs, [(r.sheet, r.row) for r in rows], issues

    start = spec.debut_document
    heads = [r for r in rows if r.code == start]
    keys = [r.key for r in heads]
    for k in {k for k in keys if keys.count(k) > 1}:
        issues.append(Issue(0, start, None, f"PIECE {k!r} utilisée par plusieurs entêtes"))
    by_key: dict[str, list[_Row]] = defaultdict(list)
    commands: list[_Row] = []
    rank = {code: i for i, code in enumerate(spec.enregistrements)}
    for r in rows:
        rec = spec.enregistrements[r.code]
        if r.code == start:
            continue
        if rec.niveau == "commande" and (r.key is None or r.key not in keys):
            commands.append(r)
        elif r.key is None or r.key not in keys:
            issues.append(Issue(0, r.code, None, f"feuille {r.sheet} ligne {r.row} : PIECE {r.key!r} sans entête {start}"))
        else:
            by_key[r.key].append(r)

    docs: list[Document] = []
    origins: list[tuple[str, int]] = []

    def emit(doc: Document, group: list[_Row]) -> None:
        # Fusionne les lignes consécutives d'un enregistrement à colonnes répétées (ED, LD).
        merged: list[tuple[_Row, list]] = []
        for r in group:
            rec = spec.enregistrements[r.code]
            if rec.repetition_depuis:
                rep = [f.nom for f in rec.champs if f.n >= rec.repetition_depuis]
                pair = tuple(r.values.get(n) for n in rep)
                if merged and merged[-1][0].code == r.code and merged[-1][0].order == r.order:
                    merged[-1][1].append(pair)
                    continue
                merged.append((r, [pair]))
            else:
                merged.append((r, None))
        for r, pairs in merged:
            values = dict(r.values)
            if pairs is not None:
                values = {"paires": pairs}
            doc.add(r.code, **values)
            origins.append((r.sheet, r.row))

    for head in heads:
        doc = Document()
        emit(doc, [head])
        group = by_key.get(head.key, [])
        header_part = sorted((r for r in group if spec.enregistrements[r.code].niveau == "entete"),
                             key=lambda r: (rank[r.code], r.row))
        line_part = sorted((r for r in group if spec.enregistrements[r.code].niveau == "ligne"),
                           key=lambda r: (r.order if r.order is not None else float("inf"),
                                          bool(spec.enregistrements[r.code].suit), r.row))
        cmd_part = sorted((r for r in group if spec.enregistrements[r.code].niveau == "commande"),
                          key=lambda r: (r.order or 0, r.row))
        emit(doc, header_part)
        emit(doc, line_part)
        emit(doc, cmd_part)
        docs.append(doc)
    for r in sorted(commands, key=lambda r: (r.order or 0, r.row)):
        doc = Document()
        emit(doc, [r])
        docs.append(doc)
    return docs, origins, issues


def build_file(spec: FormatSpec, workbook: str | Path, out: str | Path, sep: str = ";") -> list[Issue]:
    """Lit le classeur rempli, contrôle et écrit le fichier d'import (rien n'est écrit en cas d'erreur)."""
    rows = _read_rows(spec, workbook)
    docs, origins, issues = build_documents(spec, rows)
    blocked = any(i.level == "erreur" for i in issues)
    file_issues = write_file(out, spec, docs, sep, write=not blocked)
    for i in file_issues:  # renvoie l'anomalie vers la feuille et la ligne Excel d'origine
        if 0 < i.line <= len(origins):
            sheet, row = origins[i.line - 1]
            i.message = f"feuille {sheet} ligne {row} : {i.message}"
            i.line = 0
    return issues + file_issues
