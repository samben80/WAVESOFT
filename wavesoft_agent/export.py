"""Exports lisibles du catalogue : dictionnaire Markdown et classeur Excel."""

from __future__ import annotations

from pathlib import Path

from .catalog import Catalog, DbObject

KIND_LABELS = {
    "table": "Tables",
    "view": "Vues",
    "procedure": "Procédures stockées",
    "function": "Fonctions",
    "trigger": "Triggers",
    "synonym": "Synonymes",
    "sequence": "Séquences",
}


def _md_cell(value) -> str:
    return "" if value is None else str(value).replace("|", "\\|").replace("\n", " ")


def _object_md(o: DbObject) -> list[str]:
    lines = [f"### {o.full_name}", ""]
    if o.description:
        lines += [o.description, ""]
    meta = []
    if o.row_count is not None:
        meta.append(f"{o.row_count:,} lignes".replace(",", " "))
    if o.primary_key:
        meta.append("clé primaire : " + ", ".join(o.primary_key.columns))
    if o.parent:
        meta.append(f"sur {o.parent} ({', '.join(o.trigger_events)}){' — désactivé' if o.disabled else ''}")
    if o.base_object:
        meta.append(f"pointe vers {o.base_object}")
    if meta:
        lines += ["- " + m for m in meta] + [""]
    if o.columns:
        lines += ["| # | Colonne | Type | Null | Défaut | Obligatoire | Description |", "|---|---|---|---|---|---|---|"]
        for c in o.columns:
            extra = " (identité)" if c.identity else " (calculée)" if c.computed else ""
            lines.append(
                f"| {c.position} | {_md_cell(c.name)} | {c.sql_type}{extra} | {'oui' if c.nullable else 'non'} "
                f"| {_md_cell(c.default)} | {'**oui**' if c.required else ''} | {_md_cell(c.description)} |"
            )
        lines.append("")
    if o.parameters:
        lines += ["| Paramètre | Type | Sortie | Défaut |", "|---|---|---|---|"]
        for p in o.parameters:
            lines.append(f"| {p.name} | {p.data_type} | {'oui' if p.output else ''} | {'oui' if p.has_default else ''} |")
        lines.append("")
    for fk in o.foreign_keys:
        lines.append(f"- FK `{fk.name}` : ({', '.join(fk.columns)}) → {fk.ref_schema}.{fk.ref_table} ({', '.join(fk.ref_columns)})")
    if o.foreign_keys:
        lines.append("")
    if o.references:
        lines += ["Utilise : " + ", ".join(sorted(o.references)), ""]
    return lines


def to_markdown(catalog: Catalog) -> str:
    lines = [
        f"# Dictionnaire de la base {catalog.database}",
        "",
        f"- Serveur : {catalog.server}",
        f"- Version : {catalog.server_version or 'inconnue'}",
        f"- Collation : {catalog.collation or 'inconnue'}",
        f"- Extrait le : {catalog.extracted_at}",
        "",
        "| Type | Nombre |",
        "|---|---|",
    ]
    counts = catalog.counts()
    lines += [f"| {KIND_LABELS.get(k, k)} | {n} |" for k, n in counts.items()] + [""]
    for kind, label in KIND_LABELS.items():
        objs = catalog.of_kind(kind)
        if objs:
            lines += [f"## {label}", ""]
            for o in objs:
                lines += _object_md(o)
    return "\n".join(lines)


def to_excel(catalog: Catalog, path: str | Path) -> Path:
    """Classeur à trois onglets : Objets, Colonnes, Paramètres (filtrables)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    sheets = {
        "Objets": ["Schéma", "Nom", "Type", "Lignes", "Clé primaire", "Description", "Modifié le", "Utilise"],
        "Colonnes": ["Schéma", "Objet", "Type objet", "#", "Colonne", "Type", "Null", "Identité",
                     "Calculée", "Défaut", "Obligatoire", "Description"],
        "Paramètres": ["Schéma", "Objet", "#", "Paramètre", "Type", "Sortie", "Défaut"],
    }
    ws_objs = wb.active
    ws_objs.title = "Objets"
    ws = {"Objets": ws_objs}
    for name in ("Colonnes", "Paramètres"):
        ws[name] = wb.create_sheet(name)
    for name, header in sheets.items():
        ws[name].append(header)
        for cell in ws[name][1]:
            cell.font = Font(bold=True)
        ws[name].freeze_panes = "A2"

    for o in catalog.objects:
        pk = ", ".join(o.primary_key.columns) if o.primary_key else None
        ws["Objets"].append([o.schema, o.name, o.kind, o.row_count, pk, o.description, o.modified,
                             ", ".join(sorted(o.references)) or None])
        for c in o.columns:
            ws["Colonnes"].append([o.schema, o.name, o.kind, c.position, c.name, c.sql_type, c.nullable,
                                   c.identity, c.computed, c.default, c.required, c.description])
        for p in o.parameters:
            ws["Paramètres"].append([o.schema, o.name, p.position, p.name, p.data_type, p.output, p.has_default])

    for sheet in ws.values():
        sheet.auto_filter.ref = sheet.dimensions
        for col in sheet.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            sheet.column_dimensions[col[0].column_letter].width = min(max(width + 2, 8), 60)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
